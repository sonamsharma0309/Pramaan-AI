"""Safe text extraction from uploads: txt/md/csv/edi, text PDFs (with optional OCR), and docx."""
import io
import shutil

from fastapi import HTTPException, UploadFile

from .. import config

ALLOWED = {".txt", ".md", ".csv", ".edi", ".835", ".x12", ".pdf", ".docx"}
MAGIC_PDF = b"%PDF-"


def ocr_available() -> bool:
    """True when pytesseract, pdf2image and the tesseract + poppler binaries are installed."""
    try:
        import pdf2image  # noqa: F401
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return bool(shutil.which("tesseract") and shutil.which("pdftoppm"))


def _ocr_pdf(raw: bytes) -> str:
    import pytesseract
    from pdf2image import convert_from_bytes
    pages = convert_from_bytes(raw, dpi=200, first_page=1, last_page=15)
    return "\n".join(pytesseract.image_to_string(p) for p in pages)


def extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    try:
        text = "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(raw)).pages)
    except Exception:
        raise HTTPException(400, "Could not read this PDF. It may be corrupted or password protected.")
    if len(text.strip()) >= 20:
        return text
    if not ocr_available():
        raise HTTPException(400, "This PDF looks scanned (no selectable text) and OCR is not installed on this server. "
                                 "Install Tesseract + Poppler (the Docker image includes them) or paste the text instead.")
    try:
        text = _ocr_pdf(raw)
    except Exception:
        raise HTTPException(400, "OCR failed on this PDF. Try a clearer scan or paste the text.")
    if len(text.strip()) < 20:
        raise HTTPException(400, "OCR could not find readable text in this PDF.")
    return text


def extract_docx(raw: bytes) -> str:
    try:
        import docx
        d = docx.Document(io.BytesIO(raw))
    except Exception:
        raise HTTPException(400, "Could not read this Word file.")
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


async def read_text(file: UploadFile) -> str:
    raw = await file.read(config.MAX_UPLOAD_BYTES + 1)
    if len(raw) > config.MAX_UPLOAD_BYTES:
        raise HTTPException(413, f"File is larger than {config.MAX_UPLOAD_BYTES // 1048576} MB.")
    if not raw:
        raise HTTPException(400, "The file is empty.")
    name = (file.filename or "").lower()
    ext = "." + name.rsplit(".", 1)[-1] if "." in name else ""
    if ext and ext not in ALLOWED:
        raise HTTPException(415, f"Unsupported file type {ext}. Allowed: {', '.join(sorted(ALLOWED))}.")
    if raw[:5] == MAGIC_PDF or ext == ".pdf":
        if raw[:5] != MAGIC_PDF:
            raise HTTPException(415, "This file is not a valid PDF.")
        return extract_pdf(raw)
    if ext == ".docx" or raw[:2] == b"PK":
        if ext != ".docx":
            raise HTTPException(415, "Unsupported archive. Upload .docx, .pdf or text.")
        return extract_docx(raw)
    if b"\x00" in raw[:1024]:
        raise HTTPException(415, "Binary files are not supported.")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="ignore")
